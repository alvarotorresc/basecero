from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from . import contract as c
from . import seeds

N = c.PREFILLED_ROWS

def col(table, name):
    names = [x.name for x in c.TABLES[table]]
    return get_column_letter(names.index(name) + 1)

def tr(name):
    l = col("transactions", name)
    return f"transactions!${l}$2:${l}${N + 1}"

def saldo_expr(id_expr):
    e, g, i_, d, z = tr("amount"), tr("account_id"), tr("counter_account_id"), tr("type"), tr("deleted")
    return (f'VLOOKUP({id_expr},accounts!$A:$D,4,FALSE)'
            f'+SUMIFS({e},{g},{id_expr},{d},"income",{z},"<>TRUE")'
            f'+SUMIFS({e},{g},{id_expr},{d},"refund",{z},"<>TRUE")'
            f'-SUMIFS({e},{g},{id_expr},{d},"expense",{z},"<>TRUE")'
            f'-SUMIFS({e},{g},{id_expr},{d},"transfer",{z},"<>TRUE")'
            f'+SUMIFS({e},{i_},{id_expr},{d},"transfer",{z},"<>TRUE")'
            f'+SUMIFS({e},{g},{id_expr},{d},"adjustment",{z},"<>TRUE")')

def _head(ws, names, color):
    ws.sheet_properties.tabColor = color
    ws.freeze_panes = "A2"
    for j, name in enumerate(names, 1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEEEEE")

def _dv(ws, formula, letra, boolean=False):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"{letra}2:{letra}{N + 1}")

def _write_meta(wb):
    ws = wb.create_sheet("meta")
    _head(ws, ["key", "value"], c.COLOR_DATA)
    for r, (k, v) in enumerate(c.META_KEYS.items(), 2):
        ws.cell(row=r, column=1, value=k); ws.cell(row=r, column=2, value=v)
    for j, (name, values) in enumerate(c.ENUMS.items()):
        letra = get_column_letter(4 + j)
        ws.cell(row=1, column=4 + j, value=name).font = Font(bold=True)
        for r, v in enumerate(values, 2):
            ws.cell(row=r, column=4 + j, value=v)
        wb.defined_names[f"enum_{name}"] = DefinedName(f"enum_{name}",
            attr_text=f"meta!${letra}$2:${letra}${len(values) + 1}")

def _write_rows(ws, table, rows):
    names = [x.name for x in c.TABLES[table]]
    for r, row in enumerate(rows, 2):
        for j, name in enumerate(names, 1):
            if name in row:
                ws.cell(row=r, column=j, value=row[name])
        ws.cell(row=r, column=names.index("created_at") + 1, value=seeds.NOW_ISO)
        ws.cell(row=r, column=names.index("updated_at") + 1, value=seeds.NOW_ISO)
        ws.cell(row=r, column=names.index("deleted") + 1, value=False)

def _write_data_sheets(wb):
    for t in c.DATA_SHEET_ORDER:
        if t == "meta":
            continue
        ws = wb.create_sheet(t)
        names = [x.name for x in c.TABLES[t]]
        _head(ws, names, c.COLOR_DATA)
        for x in c.TABLES[t]:                     # validaciones por tipo de columna
            letra = col(t, x.name)
            if x.kind.startswith("enum:"):
                _dv(ws, f"=enum_{x.kind.split(':')[1]}", letra)
            elif x.kind == "bool":
                _dv(ws, '"TRUE,FALSE"', letra)
        for x in c.TABLES[t]:                     # system y _picker ocultas
            if x.name in c.SYSTEM_COLUMNS or x.name == "_picker":
                ws.column_dimensions[col(t, x.name)].hidden = True
    _write_rows(wb["accounts"], "accounts", seeds.SEED_ACCOUNTS)
    _write_rows(wb["categories"], "categories", seeds.SEED_CATEGORIES)
    ws = wb["categories"]                          # _picker: elegible si es hoja
    for r in range(2, 202):
        ws.cell(row=r, column=8, value=(
            f'=IF($A{r}="","",IF(COUNTIF($C:$C,$A{r})>0,"",'
            f'IF($C{r}="",$B{r},VLOOKUP($C{r},$A:$B,2,FALSE)&" → "&$B{r})))'))
    ws = wb["transactions"]                        # _my_amount y _sara_amount
    for r in range(2, N + 2):
        ws.cell(row=r, column=16, value=(
            f'=IF($E{r}="",0,IF($N{r}=TRUE,ROUND($E{r}*IF($O{r}="",'
            f'IFERROR(VLOOKUP($C{r},periods!$A:$F,6,FALSE),100),$O{r})/100,2),$E{r}))'))
        ws.cell(row=r, column=17, value=f'=IF($N{r}=TRUE,ROUND($E{r}-$P{r},2),0)')
    _dv(ws, "=accounts_names", col("transactions", "_account"))
    _dv(ws, "=accounts_names", col("transactions", "_counter_account"))
    _dv(ws, "=categories_pickers", col("transactions", "_category"))
    _dv(ws, "=rules_names", col("transactions", "_rule"))
    for nombre, ref in [("accounts_names", "accounts!$B$2:$B$100"),
                        ("categories_pickers", "categories!$H$2:$H$201"),
                        ("rules_names", "recurring_rules!$B$2:$B$100"),
                        ("periods_names", "periods!$B$2:$B$100")]:
        wb.defined_names[nombre] = DefinedName(nombre, attr_text=ref)

def build(path):
    wb = Workbook()
    wb.remove(wb.active)
    _write_meta(wb)
    _write_data_sheets(wb)
    from .dashboards import write_dashboards   # Task 4; hasta entonces, stub
    write_dashboards(wb)
    wb.save(path)
