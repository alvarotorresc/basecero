from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from . import contract as c
from . import seeds
from .build_xlsx import col, tr, saldo_expr

B = Font(bold=True)

def _sheet(wb, name):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = c.COLOR_DASH
    ws["A1"].font = B
    return ws

def _gasto(cat_id_expr, period_cell):
    p, e, d, z, k, s, cc = (tr("_my_amount"), tr("amount"), tr("type"), tr("deleted"),
                            tr("category_id"), tr("ref_id"), tr("period_id"))
    base = f'({cc}={period_cell})*({z}=FALSE)'
    filtro = f'({k}={cat_id_expr})*' if cat_id_expr else ""
    return (f'SUMPRODUCT({filtro}{base}*({d}="expense")*{p})'
            f'-SUMPRODUCT({filtro}{base}*({d}="refund")*({s}="")*{e})')

def _resumen(wb):
    ws = _sheet(wb, "Resumen del periodo")
    ws["A1"] = "Periodo:"
    dv = DataValidation(type="list", formula1="=periods_names", allow_blank=True)
    ws.add_data_validation(dv); dv.add("B1")
    ws["B2"] = '=IFERROR(INDEX(periods!$A:$A,MATCH($B$1,periods!$B:$B,0)),"")'
    ws["B3"] = '=IFERROR(INDEX(periods!$A:$A,MATCH($B$1,periods!$B:$B,0)-1),"")'
    e, d, z, cc, n, q, r2, v, k = (tr("amount"), tr("type"), tr("deleted"), tr("period_id"),
        tr("is_shared"), tr("_sara_amount"), tr("settled"), tr("external_id"), tr("category_id"))
    ws["A5"], ws["B5"] = "Ingresos", (f'=SUMIFS({e},{cc},$B$2,{d},"income",{z},"<>TRUE")')
    ws["A6"], ws["B6"] = "Gasto (mi parte)", "=" + _gasto("", "$B$2")
    ws["A7"], ws["B7"] = "Ahorro €", "=B5-B6"
    ws["A8"], ws["B8"] = "Tasa de ahorro", '=IF(B5=0,"",B7/B5)'
    comp = f'({cc}=$B$2)*({n}=TRUE)*({d}="expense")*({z}=FALSE)'
    ws["A10"], ws["B10"] = "Compartido total", f'=SUMPRODUCT({comp}*{e})'
    ws["A11"], ws["B11"] = "Mi parte", f'=SUMPRODUCT({comp}*{tr("_my_amount")})'
    ws["A12"], ws["B12"] = "Parte de Sara", f'=SUMPRODUCT({comp}*{q})'
    ws["A13"], ws["B13"] = "Pendiente de cobro (global)", \
        f'=SUMPRODUCT(({n}=TRUE)*({r2}=FALSE)*({d}="expense")*({z}=FALSE)*{q})'
    ws["A15"], ws["B15"] = "Sin categorizar", \
        f'=SUMPRODUCT(({v}<>"")*({k}="")*({z}=FALSE)*1)'
    for j, t in enumerate(["Categoría", "Subcategoría", "Este periodo", "Anterior", "Δ"], 1):
        ws.cell(row=17, column=j, value=t).font = B
    by_id = {x["id"]: x for x in seeds.SEED_CATEGORIES}
    hojas = [x for x in seeds.SEED_CATEGORIES if x["flow"] == "expense"
             and not any(y["parent_id"] == x["id"] for y in seeds.SEED_CATEGORIES)]
    for i, cat in enumerate(hojas):
        r = 18 + i
        raiz = by_id[cat["parent_id"]]["name"] if cat["parent_id"] else cat["name"]
        ws.cell(row=r, column=1, value=raiz)
        ws.cell(row=r, column=2, value=cat["name"])
        ws.cell(row=r, column=3, value="=" + _gasto(f'"{cat["id"]}"', "$B$2"))
        ws.cell(row=r, column=4, value="=" + _gasto(f'"{cat["id"]}"', "$B$3"))
        ws.cell(row=r, column=5, value=f'=IF(OR(D{r}="",D{r}=0),"",C{r}/D{r}-1)')

def _patrimonio(wb):
    ws = _sheet(wb, "Patrimonio y objetivos")
    ws["A1"] = "Cuentas"
    for i in range(12):                      # 12 cuentas de plantilla
        r, ar = 3 + i, 2 + i
        ws.cell(row=r, column=1, value=f'=IF(accounts!$B{ar}="","",accounts!$B{ar})')
        ws.cell(row=r, column=2, value=(
            f'=IF(accounts!$B{ar}="","",{saldo_expr(f"accounts!$A{ar}")})'))
    ws["A16"], ws["B16"] = "PATRIMONIO NETO", "=SUM(B3:B14)"
    ws["A16"].font = B
    ws["A18"] = "Evolución por periodo"
    b, e, d, z = tr("date"), tr("amount"), tr("type"), tr("deleted")
    for i in range(12):                      # 12 periodos de serie
        r, pr = 19 + i, 2 + i
        ws.cell(row=r, column=1, value=f'=IF(periods!$B{pr}="","",periods!$B{pr})')
        ws.cell(row=r, column=2, value=(
            f'=IF(periods!$D{pr}="","",SUM(accounts!$D$2:$D$100)'
            f'+SUMPRODUCT(({b}<=periods!$D{pr})*(({d}="income")+({d}="refund")+({d}="adjustment"))*({z}=FALSE)*{e})'
            f'-SUMPRODUCT(({b}<=periods!$D{pr})*({d}="expense")*({z}=FALSE)*{e}))'))
    ws["A32"], ws["B32"] = "Periodo abierto", \
        '=IFERROR(INDEX(periods!$A:$A,MATCH("open",periods!$E:$E,0)),"")'
    ws["A33"] = "Objetivos"
    gasto_medio = (f'((SUMPRODUCT(({tr("period_id")}<>$B$32)*({tr("type")}="expense")*({tr("deleted")}=FALSE)*{tr("_my_amount")})'
                   f'-SUMPRODUCT(({tr("period_id")}<>$B$32)*({tr("type")}="refund")*({tr("ref_id")}="")*({tr("deleted")}=FALSE)*{tr("amount")}))'
                   f'/MAX(1,COUNTIF(periods!$E:$E,"closed")))')
    for j, t in enumerate(["Objetivo", "Tipo", "Progreso", "Barra"], 1):
        ws.cell(row=34, column=j, value=t).font = B
    for i in range(10):                      # 10 goals de plantilla
        r, gr = 35 + i, 2 + i
        saldo = saldo_expr(f"goals!$H{gr}")
        gasto_cap = _gasto(f"goals!$I{gr}", "$B$32")
        ws.cell(row=r, column=1, value=f'=IF(goals!$B{gr}="","",goals!$B{gr})')
        ws.cell(row=r, column=2, value=f'=IF(goals!$B{gr}="","",goals!$C{gr})')
        ws.cell(row=r, column=3, value=(
            f'=IF(goals!$B{gr}="","",IFERROR(IFS('
            f'goals!$C{gr}="savings_target",({saldo})/goals!$D{gr},'
            f'goals!$C{gr}="provision",({saldo})/goals!$D{gr},'
            f'goals!$C{gr}="emergency_fund",({saldo})/(goals!$E{gr}*{gasto_medio}),'
            f'goals!$C{gr}="spending_cap",({gasto_cap})/goals!$D{gr},'
            f'goals!$C{gr}="savings_rate",1),""))'))
        ws.cell(row=r, column=4, value=(
            f'=IF(C{r}="","",REPT("█",MIN(10,ROUND(C{r}*10,0)))&" "&TEXT(C{r},"0%"))'))
    # savings_rate=1 es un marcador honesto del MVP: la tasa objetivo se compara a ojo
    # con la tasa del Resumen; la app lo calculará de verdad. Documentado en README.

def _prevision(wb):
    ws = _sheet(wb, "Previsión")
    ws["A1"], ws["B1"] = "Periodo abierto", \
        '=IFERROR(INDEX(periods!$A:$A,MATCH("open",periods!$E:$E,0)),"")'
    ws["A2"], ws["B2"] = "Mi % este periodo", \
        '=IFERROR(VLOOKUP($B$1,periods!$A:$F,6,FALSE),100)'
    ws["C1"], ws["D1"] = "Mes del periodo", \
        '=IFERROR(MONTH(VLOOKUP($B$1,periods!$A:$C,3,FALSE)+15),MONTH(TODAY()))'
    for j, t in enumerate(["Regla", "Frecuencia", "¿Aplica este mes?", "Mi importe", "Tipo", "Estado"], 1):
        ws.cell(row=3, column=j, value=t).font = B
    u, cc, k, e = tr("rule_id"), tr("period_id"), tr("category_id"), tr("amount")
    for i in range(20):                      # 20 reglas de plantilla
        r, rr = 4 + i, 2 + i
        ws.cell(row=r, column=1, value=f'=IF(recurring_rules!$B{rr}="","",recurring_rules!$B{rr})')
        ws.cell(row=r, column=2, value=f'=IF(recurring_rules!$B{rr}="","",recurring_rules!$H{rr})')
        ws.cell(row=r, column=3, value=(
            f'=IF(recurring_rules!$B{rr}="","",IF(recurring_rules!$L{rr}=FALSE,"no",'
            f'IF(OR(recurring_rules!$H{rr}="monthly",recurring_rules!$H{rr}="weekly"),"sí",'
            f'IF(recurring_rules!$J{rr}="","no",'
            f'IF(recurring_rules!$H{rr}="yearly",IF(recurring_rules!$J{rr}=$D$1,"sí","no"),'
            f'IF(MOD($D$1-recurring_rules!$J{rr},3)=0,"sí","no"))))))'))
        ws.cell(row=r, column=4, value=(
            f'=IF(recurring_rules!$B{rr}="","",IF(recurring_rules!$K{rr}=TRUE,'
            f'ROUND(recurring_rules!$D{rr}*$B$2/100,2),recurring_rules!$D{rr}))'))
        ws.cell(row=r, column=5, value=f'=IF(recurring_rules!$B{rr}="","",recurring_rules!$C{rr})')
        ws.cell(row=r, column=6, value=(
            f'=IF(recurring_rules!$B{rr}="","",'
            f'IF(OR(COUNTIFS({u},recurring_rules!$A{rr},{cc},$B$1)>0,'
            f'COUNTIFS({k},recurring_rules!$E{rr},{e},recurring_rules!$D{rr},{cc},$B$1)>0),'
            f'"✅ pagado","⏳ pendiente"))'))
    ws["A26"], ws["B26"] = "Comprometido restante", \
        '=SUMPRODUCT(($C$4:$C$23="sí")*($F$4:$F$23="⏳ pendiente")*($E$4:$E$23<>"income")*($D$4:$D$23))'
    n, q, r2, d, z = tr("is_shared"), tr("_sara_amount"), tr("settled"), tr("type"), tr("deleted")
    ws["A27"], ws["B27"] = "Pendiente de Sara", \
        f'=SUMPRODUCT(({n}=TRUE)*({r2}=FALSE)*({d}="expense")*({z}=FALSE)*{q})'
    ws["A28"], ws["B28"] = "Saldo N26", "=" + saldo_expr('"acc-n26"')
    ws["A29"], ws["B29"] = "Disponible real", "=B28-B26+B27"
    ws["A29"].font = B

def write_dashboards(wb):
    _resumen(wb); _patrimonio(wb); _prevision(wb)
