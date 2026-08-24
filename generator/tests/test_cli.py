import subprocess, sys
from openpyxl import load_workbook

def test_cli_genera_el_archivo(tmp_path):
    out = tmp_path / "BaseCero.xlsx"
    res = subprocess.run([sys.executable, "-m", "basecero_generator.cli", str(out)],
                         capture_output=True, text=True)
    assert res.returncode == 0 and out.exists()
    wb = load_workbook(str(out))
    assert len(wb.sheetnames) == 11
