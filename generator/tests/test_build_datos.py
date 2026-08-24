import pytest
from openpyxl import load_workbook
from basecero_generator import contract as c
from basecero_generator.build_xlsx import build, col

@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    p = tmp_path_factory.mktemp("out") / "BaseCero.xlsx"
    build(str(p))
    return load_workbook(str(p))

def test_pestanas_de_datos_en_orden(wb):
    assert wb.sheetnames[:8] == c.DATA_SHEET_ORDER
    for name in c.DATA_SHEET_ORDER:
        assert wb[name].sheet_properties.tabColor.rgb.endswith(c.COLOR_DATA)

def test_cabeceras_exactas(wb):
    for t, columns in c.TABLES.items():
        fila1 = [cell.value for cell in wb[t][1][: len(columns)]]
        assert fila1 == [x.name for x in columns], t

def test_meta(wb):
    ws = wb["meta"]
    assert (ws["A1"].value, ws["B1"].value) == ("key", "value")
    keys = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(2, 5)}
    assert keys == c.META_KEYS
    cabeceras_enums = [ws.cell(row=1, column=4 + i).value for i in range(len(c.ENUMS))]
    assert set(cabeceras_enums) == set(c.ENUMS)

def test_semillas_escritas(wb):
    assert wb["accounts"]["B2"].value == "N26"
    assert wb["accounts"]["C5"].value == "liability"
    nombres = [wb["categories"].cell(row=r, column=2).value for r in range(2, 60)]
    assert "Supermercado" in nombres and "Nómina" in nombres

def test_columnas_system_ocultas_y_freeze(wb):
    for t, columns in c.TABLES.items():
        ws = wb[t]
        assert ws.freeze_panes == "A2"
        for name in c.SYSTEM_COLUMNS:
            assert ws.column_dimensions[col(t, name)].hidden, (t, name)

def test_formulas_precargadas_en_transactions(wb):
    ws = wb["transactions"]
    assert str(ws["P2"].value).startswith("=IF(")
    assert str(ws[f"Q{c.PREFILLED_ROWS + 1}"].value).startswith("=IF(")

def test_validaciones_presentes(wb):
    ws = wb["transactions"]
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert "=enum_transaction_types" in formulas
    assert "=accounts_names" in formulas and "=categories_pickers" in formulas
