import pytest
from openpyxl import load_workbook
from basecero_generator.build_xlsx import build
from basecero_generator import contract as c

@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    p = tmp_path_factory.mktemp("out") / "BaseCero.xlsx"
    build(str(p))
    return load_workbook(str(p))

def test_orden_y_color(wb):
    assert wb.sheetnames == c.DATA_SHEET_ORDER + \
        ["Resumen del periodo", "Patrimonio y objetivos", "Previsión"]
    for name in wb.sheetnames[8:]:
        assert wb[name].sheet_properties.tabColor.rgb.endswith(c.COLOR_DASH)

def test_resumen(wb):
    ws = wb["Resumen del periodo"]
    assert "INDEX(periods" in str(ws["B2"].value)
    assert "SUMIFS" in str(ws["B5"].value)          # ingresos
    assert "_" not in str(ws["B6"].value) and "SUMPRODUCT" in str(ws["B6"].value)
    assert "SUMPRODUCT" in str(ws["B13"].value)     # pendiente Sara
    fila_super = [r for r in range(18, 60) if ws.cell(row=r, column=2).value == "Supermercado"]
    assert fila_super and "cat-alimentacion-supermercado" in str(
        ws.cell(row=fila_super[0], column=3).value)

def test_patrimonio(wb):
    ws = wb["Patrimonio y objetivos"]
    assert "VLOOKUP(\"acc-n26\"" in str(ws["B3"].value)
    assert str(ws["B7"].value).startswith("=SUM(")   # patrimonio neto
    assert "REPT" in str(ws["D31"].value)            # barra de objetivos, fila 1 de plantilla

def test_prevision(wb):
    ws = wb["Previsión"]
    assert 'MATCH("open"' in str(ws["B1"].value)
    assert "COUNTIFS" in str(ws["F4"].value)         # estado pagado/pendiente
    assert "Disponible real" in [ws.cell(row=r, column=1).value for r in range(25, 32)]
